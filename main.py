from xml.dom import minidom
import openpyxl

book = openpyxl.load_workbook("Price.xlsx", read_only=True)
sheet = book.active


root = minidom.Document()
xml = root.createElement('ARTICLES')
root.appendChild(xml)

for row in range(1, sheet.max_row):
    if (sheet[row][0].value == None):
        break
    productChild = root.createElement('Art')
    productChild.setAttribute('SERIAL', '')
    productChild.setAttribute('CODE', str(int(sheet[row][0].value)))
    productChild.setAttribute('NAME', str(sheet[row][1].value))
    productChild.setAttribute('FULLNAME', '')
    productChild.setAttribute('PRICE', str(sheet[row][2].value))
    productChild.setAttribute('AMOUNT', '0')
    productChild.setAttribute('NUMDEP', '1')
    productChild.setAttribute('AGROUP', '1')
    productChild.setAttribute('TAX', '5')
    productChild.setAttribute('FLAG', '0')
    productChild.setAttribute('AC0', '0')
    productChild.setAttribute('OPERATION', '')
    productChild.setAttribute('RESULT', '')
    productChild.setAttribute('LCODE', '')
    productChild.setAttribute('LAMOUNT', '')
    productChild.setAttribute('BARCODE', '')
    productChild.setAttribute('BARCODE2', '')
    productChild.setAttribute('BARCODE3', '')
    productChild.setAttribute('BARCODE4', '')
    productChild.setAttribute('RESULTTEXT', '')
    productChild.setAttribute('CUSTOMCODE', '')
    productChild.setAttribute('FRACTIONALQTY', '1')
    productChild.setAttribute('UNIT', '0')

    xml.appendChild(productChild)

    xml_str = root.toprettyxml(indent="\t")

    save_path_file = "FileConvert.xml"


with open(save_path_file, "w", encoding='utf-8') as f:
    f.write(xml_str)
